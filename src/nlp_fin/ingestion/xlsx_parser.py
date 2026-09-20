"""Парсер книг Excel."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from nlp_fin.ingestion.base import ReportParser
from nlp_fin.ingestion.common import rows_to_text
from nlp_fin.ingestion.models import (
    BlockType,
    DocumentType,
    ParsedDocument,
    Section,
    SourceInfo,
)


class XlsxParser(ReportParser):
    """Извлекает содержимое листов книги Excel в виде таблиц."""

    def parse(self, path: Path) -> ParsedDocument:
        sections: list[Section] = []
        raw_parts: list[str] = []

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                sections.append(Section(page=1, block_type=BlockType.TITLE, text=sheet.title))
                raw_parts.append(sheet.title)

                rows: list[list[str]] = []
                for row in sheet.iter_rows(values_only=True):
                    cells = ["" if cell is None else str(cell).strip() for cell in row]
                    if any(cells):
                        rows.append(cells)

                if rows:
                    sections.append(
                        Section(
                            page=1,
                            block_type=BlockType.TABLE,
                            text=rows_to_text(rows),
                            rows=rows,
                        )
                    )
                    raw_parts.append(rows_to_text(rows))
        finally:
            workbook.close()

        return ParsedDocument(
            source=SourceInfo(filename=path.name, document_type=DocumentType.XLSX),
            sections=sections,
            raw_text="\n".join(raw_parts),
        )