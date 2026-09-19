"""Тесты парсеров документов."""

from __future__ import annotations

from pathlib import Path

import pymupdf
import pytesseract
import pytest
from docx import Document as WordDocument
from openpyxl import Workbook

from nlp_fin.ingestion.models import BlockType, DocumentType
from nlp_fin.ingestion.registry import detect_document_type, parse_document

TABLE_MARKER = "Выручка"
WINDOWS_FONT = "notos"


def make_text_pdf(path: Path) -> None:
    """Создаёт PDF с текстовым слоем, включающий таблицу, шапку и подвал."""
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 25), "Шапка отчёта", fontname=WINDOWS_FONT, fontsize=8)
    page.insert_text((72, 100), "Годовой отчёт за 2023 год", fontname=WINDOWS_FONT, fontsize=16)
    page.insert_textbox(
        pymupdf.Rect(72, 130, 523, 180),
        "Выручка компании за отчётный период составила 123,4 млн руб.",
        fontname=WINDOWS_FONT,
        fontsize=10,
    )
    _draw_table(page, [["Показатель", "Значение"], [TABLE_MARKER, "123 456"]], 72, 200)
    page.insert_text((72, 820), "Страница 1", fontname=WINDOWS_FONT, fontsize=8)

    second = doc.new_page(width=595, height=842)
    second.insert_textbox(
        pymupdf.Rect(72, 60, 523, 200),
        "Во втором разделе приведены пояснения к бухгалтерской отчётности.",
        fontname=WINDOWS_FONT,
        fontsize=10,
    )

    doc.save(path)


def make_scan_pdf(path: Path) -> None:
    """Создаёт PDF из растрового изображения страницы."""
    source = pymupdf.open()
    text_page = source.new_page()
    text_page.insert_text(
        (72, 100), "Выручка компании за 2023 год", fontname=WINDOWS_FONT, fontsize=12
    )
    pix = text_page.get_pixmap(dpi=150, alpha=False)

    output = pymupdf.open()
    image_page = output.new_page(width=595, height=842)
    image_page.insert_image(pymupdf.Rect(0, 0, 595, 842), stream=pix.tobytes("png"))
    output.save(path)


def make_xlsx(path: Path) -> None:
    """Создаёт книгу Excel с одним листом."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Отчёт о прибылях"
    sheet.append(["Показатель", "2023"])
    sheet.append([TABLE_MARKER, 123456.0])
    workbook.save(path)


def make_docx(path: Path) -> None:
    """Создаёт документ Word с заголовком, абзацем и таблицей."""
    document = WordDocument()
    document.add_heading("Годовой отчёт", level=1)
    document.add_paragraph("Выручка компании составила 123,4 млн руб.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Показатель"
    table.cell(0, 1).text = "Значение"
    table.cell(1, 0).text = TABLE_MARKER
    table.cell(1, 1).text = "123 456"
    document.save(path)


def make_html(path: Path) -> None:
    """Создаёт HTML-документ с заголовком, абзацем и таблицей."""
    content = """<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>Отчёт</title></head>
<body>
  <h1>Годовой отчёт</h1>
  <p>Выручка компании составила 123,4 млн руб.</p>
  <table>
    <tr><th>Показатель</th><th>Значение</th></tr>
    <tr><td>Выручка</td><td>123 456</td></tr>
  </table>
</body></html>"""
    path.write_text(content, encoding="utf-8")


def _draw_table(page: pymupdf.Page, rows: list[list[str]], x0: float, y0: float) -> None:
    """Рисует таблицу линиями на странице PDF."""
    col_widths = [230, 221]
    row_height = 30

    x_edges = [x0]
    for width in col_widths:
        x_edges.append(x_edges[-1] + width)

    for index in range(len(rows) + 1):
        y = y0 + index * row_height
        page.draw_line(pymupdf.Point(x_edges[0], y), pymupdf.Point(x_edges[-1], y))
    for x in x_edges:
        page.draw_line(pymupdf.Point(x, y0), pymupdf.Point(x, y0 + len(rows) * row_height))

    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            page.insert_text(
                (x_edges[j] + 4, y0 + i * row_height + 20),
                cell,
                fontname=WINDOWS_FONT,
                fontsize=10,
            )


def test_detect_text_pdf(tmp_path: Path) -> None:
    path = tmp_path / "report.pdf"
    make_text_pdf(path)
    assert detect_document_type(path) is DocumentType.PDF_TEXT


def test_detect_scan_pdf(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    make_scan_pdf(path)
    assert detect_document_type(path) is DocumentType.PDF_SCAN


@pytest.mark.parametrize(
    ("extension", "expected"),
    [
        (".xlsx", DocumentType.XLSX),
        (".docx", DocumentType.DOCX),
        (".html", DocumentType.HTML),
    ],
)
def test_detect_document_type(
    tmp_path: Path, extension: str, expected: DocumentType
) -> None:
    path = tmp_path / f"report{extension}"
    path.write_bytes(b"")
    assert detect_document_type(path) is expected


def test_detect_unknown_extension(tmp_path: Path) -> None:
    path = tmp_path / "report.doc"
    path.write_bytes(b"")
    with pytest.raises(ValueError):
        detect_document_type(path)


def test_parse_text_pdf(tmp_path: Path) -> None:
    path = tmp_path / "report.pdf"
    make_text_pdf(path)
    document = parse_document(path)

    assert document.source.document_type is DocumentType.PDF_TEXT
    assert document.raw_text
    block_types = {section.block_type for section in document.sections}
    assert BlockType.TABLE in block_types
    assert BlockType.HEADER in block_types
    assert BlockType.FOOTER in block_types
    table = next(section for section in document.sections if section.block_type is BlockType.TABLE)
    assert table.rows[1][0] == TABLE_MARKER


def test_parse_scan_pdf(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "scan.pdf"
    make_scan_pdf(path)
    monkeypatch.setattr(
        pytesseract,
        "image_to_string",
        lambda image, lang=None, config=None: "Выручка компании за 2023 год",
    )
    document = parse_document(path)

    assert document.source.document_type is DocumentType.PDF_SCAN
    assert any(TABLE_MARKER in section.text for section in document.sections)


def test_parse_xlsx(tmp_path: Path) -> None:
    path = tmp_path / "report.xlsx"
    make_xlsx(path)
    document = parse_document(path)

    assert document.source.document_type is DocumentType.XLSX
    titles = [
        section for section in document.sections if section.block_type is BlockType.TITLE
    ]
    assert any(section.text == "Отчёт о прибылях" for section in titles)
    table = next(section for section in document.sections if section.block_type is BlockType.TABLE)
    assert table.rows[1][0] == TABLE_MARKER


def test_parse_docx(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    make_docx(path)
    document = parse_document(path)

    assert document.source.document_type is DocumentType.DOCX
    assert document.sections[0].block_type is BlockType.TITLE
    block_types = {section.block_type for section in document.sections}
    assert BlockType.TABLE in block_types
    assert any(section.text and TABLE_MARKER in section.text for section in document.sections)


def test_parse_html(tmp_path: Path) -> None:
    path = tmp_path / "report.html"
    make_html(path)
    document = parse_document(path)

    assert document.source.document_type is DocumentType.HTML
    assert document.sections[0].block_type is BlockType.TITLE
    block_types = {section.block_type for section in document.sections}
    assert BlockType.TABLE in block_types
    assert any(section.text and TABLE_MARKER in section.text for section in document.sections)